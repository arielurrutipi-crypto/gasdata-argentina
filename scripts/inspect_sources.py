"""Read-only checks of public inputs for GasData's automatic readers."""
import io,json,re,zipfile,urllib.request,urllib.parse
from concurrent.futures import ThreadPoolExecutor
from bs4 import BeautifulSoup
import openpyxl,xlrd,pdfplumber

def fetch(url):
    with urllib.request.urlopen(urllib.request.Request(url,headers={"User-Agent":"Mozilla/5.0 GasDataArgentina source validation"}),timeout=45) as r:
        return r.read(100_000_000)

def inspect(kind,url):
    try:
        raw=fetch(url)
        if kind=='json': return json.loads(raw)
        if kind=='xls':
            book=xlrd.open_workbook(file_contents=raw)
            return [{'name':s.name,'head':[s.row_values(i) for i in range(min(8,s.nrows))],'tail':[s.row_values(i) for i in range(max(0,s.nrows-3),s.nrows)]} for s in book.sheets()]
        if kind=='xlsx':
            book=openpyxl.load_workbook(io.BytesIO(raw),read_only=True,data_only=True)
            return [{'name':s.title,'head':list(s.values)[:9],'tail':list(s.values)[-3:]} for s in book]
        if kind=='zip':
            z=zipfile.ZipFile(io.BytesIO(raw)); out={'names':z.namelist(),'bytes':len(raw)}
            files=[]
            for name in z.namelist():
                if name.lower().endswith(('.xls','.xlsx','.csv')):
                    data=z.read(name)
                    if name.lower().endswith('.xls'):
                        book=xlrd.open_workbook(file_contents=data)
                        files.append({'name':name,'sheets':[{'name':s.name,'rows':s.nrows,'cols':s.ncols,'head':[s.row_values(i) for i in range(min(8,s.nrows))],'tail':[s.row_values(i) for i in range(max(0,s.nrows-3),s.nrows)]} for s in book.sheets()]})
                    elif name.lower().endswith('.xlsx'):
                        book=openpyxl.load_workbook(io.BytesIO(data),read_only=True,data_only=True)
                        files.append({'name':name,'sheets':[{'name':s.title,'rows':s.max_row,'cols':s.max_column,'head':list(s.values)[:8]} for s in book]})
                    else: files.append({'name':name,'head':data[:4500].decode('utf-8','replace')})
                    if len(files)>=3: break
            out['files']=files
            return out
        if kind=='pdf':
            with pdfplumber.open(io.BytesIO(raw)) as pdf:
                return {'pages':len(pdf.pages),'text':'\n'.join(p.extract_text() or '' for p in pdf.pages)[:24000],
                        'words':[{k:w[k] for k in ('text','x0','top','x1','bottom')} for w in pdf.pages[0].extract_words() if re.search(r'\d|TGN|TGS',w['text'])][-130:]}
        soup=BeautifulSoup(raw,'html.parser')
        links=[{'text':a.get_text(' ',strip=True),'url':urllib.parse.urljoin(url,a['href'])} for a in soup.find_all('a',href=True)]
        out={'bytes':len(raw),'text':soup.get_text(' ',strip=True)[-11000:],'links':[x for x in links if re.search(r'csv|xls|zip|pdf|descarg|paridad|precio|sesco|producc|page=|aviso|detalle|normas',x['url']+' '+x['text'],re.I)][:60]}
        if 'boletinoficial.gba' in url:
            out['html']=str(soup)[str(soup).lower().find('fecha de publicaci')-1200:][:14000]
        return out
    except Exception as e:return {'error':str(e)}

TASKS={
 'eia_xls':('xls','https://www.eia.gov/dnav/ng/hist_xls/RNGWHHDd.xls'),
 'wb_xlsx':('xlsx','https://thedocs.worldbank.org/en/doc/74e8be41ceb20fa0da750cda2f6b9e4e-0050012026/related/CMO-Historical-Data-Monthly.xlsx'),
 'glp_api':('json','https://datos.energia.gob.ar/api/3/action/package_show?id=glp'),
 'gn_api':('json','https://datos.energia.gob.ar/api/3/action/package_show?id=precios-de-gas-natural'),
 'bagsa_gn':('pdf','https://www.bagsa.com.ar/wp-content/uploads/2026/09/NATURGY-GN-539-26.pdf'),
 'bagsa_glp':('pdf','https://www.bagsa.com.ar/wp-content/uploads/2026/09/NATURGY-GLP-539-26.pdf'),
 'production':('zip','https://www.energia.gob.ar/contenidos/archivos/Reorganizacion/informacion_del_mercado/mercado_hidrocarburos/tablas_dinamicas/upstream/sescoweb_produccion.zip'),
 'monthly':('pdf','https://www.enargas.gob.ar/secciones/transporte-y-distribucion/datos-operativos-despacho/graficos-programacion/9/PEI_202607.pdf'),
 'linepack':('pdf','https://www.enargas.gob.ar/secciones/transporte-y-distribucion/datos-operativos-despacho/graficos-programacion/5/LPG_20260827.pdf'),
 'glp':('html','https://www.argentina.gob.ar/economia/energia/hidrocarburos/gas-licuado-de-petroleo/precios-y-volumenes'),
 'worldbank':('html','https://www.worldbank.org/en/research/commodity-markets'),
 'bopba':('html','https://boletinoficial.gba.gob.ar/buscar?'+urllib.parse.urlencode({'search[date_gteq]':'01/01/2026','search[section]':'OFICIAL','search[words]':'BAGSA','search[sort]':'by_date_desc','commit':'Buscar'})),
 'bagsa':('html','https://www.bagsa.com.ar/index.php/tarifas/'),
 'tariff':('html','https://www.enargas.gob.ar/secciones/precios-y-tarifas/resoluciones-tarifas-vigentes.php'),
 'henry':('html','https://www.eia.gov/dnav/ng/hist/rngwhhdD.htm'),
 'propane':('html','https://www.eia.gov/dnav/pet/hist/EER_EPLLPA_PF4_Y44MB_DPGD.htm')
}
if __name__=='__main__':
    with ThreadPoolExecutor(max_workers=6) as ex:
        futures={name:ex.submit(inspect,*args) for name,args in TASKS.items()}
        for name,f in futures.items():print('SOURCE_REPORT '+json.dumps({'name':name,'result':f.result()},ensure_ascii=False,default=str),flush=True)
